from __future__ import annotations

from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.schemas.bom import PricedBOM


def export_excel(priced_bom: PricedBOM) -> bytes:
    """Generate a simple XLSX representation of the priced BOM.

    The implementation purposely avoids heavy dependencies like pandas so
    that the service can operate in minimal environments (e.g. CI).  We build
    the workbook using :mod:`openpyxl` directly which is lightweight and
    sufficient for the tabular data used in tests.
    """

    wb = Workbook()
    ws = wb.active

    headers = list(priced_bom.items[0].model_dump().keys()) if priced_bom.items else []
    ws.append(headers)
    for item in priced_bom.items:
        row = [item.model_dump().get(h) for h in headers]
        ws.append(row)

    # autosize columns for neatness (not essential but nice for manual checks)
    for i, h in enumerate(headers, start=1):
        column = get_column_letter(i)
        max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in ws[column])
        ws.column_dimensions[column].width = max(10, min(50, max_len + 2))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_json(priced_bom: PricedBOM) -> bytes:
    return json.dumps(priced_bom.model_dump(), ensure_ascii=False).encode("utf-8")


def report_pdf(priced_bom: PricedBOM, title: str) -> bytes:
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4
    y = height - 40
    c.setFont("Helvetica", 14)
    c.drawString(40, y, title)
    c.setFont("Helvetica", 10)
    y -= 30
    for it in priced_bom.items:
        line = f"{it.name} ({it.unit}): {it.qty} x {it.unit_price} {it.currency}"
        c.drawString(40, y, line)
        y -= 15
        if y < 40:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = height - 40
    c.save()
    return buf.getvalue()
